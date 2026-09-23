"""Genera los documentos descargables: facturas y reportes de ventas.

Aquí vive todo lo que produce un archivo (PDF con ReportLab, hoja de
cálculo con openpyxl), separado de las rutas. Las rutas se ocupan de
permisos y de consultar la base de datos; este módulo solo recibe datos
ya listos y devuelve los bytes del archivo.

Los documentos se arman con la identidad de la tienda: cabecera negra,
acento rojo y la información ordenada en tablas.
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# --- Identidad de la tienda ---
NEGRO = colors.HexColor("#000000")
ROJO = colors.HexColor("#FF003D")
GRIS = colors.HexColor("#6B7280")
GRIS_SUAVE = colors.HexColor("#F3F4F6")
BORDE = colors.HexColor("#E3E3E8")

EMPRESA = {
    "nombre": "PhoneStore",
    "descripcion": "Celulares, accesorios y servicio técnico",
    "direccion": "Calle 45 #12-34, Medellín, Antioquia",
    "telefono": "+57 314 772 8502",
    "email": "contacto@phonestore.com",
    "nit": "901.456.789-0",
}

# En la base se guardan como claves; en el documento se leen como texto.
METODOS_PAGO = {
    "contraentrega": "Contra entrega",
    "transferencia": "Transferencia",
    "efectivo": "Efectivo",
}


def _metodo_pago(clave) -> str:
    return METODOS_PAGO.get(clave or "", (clave or "No registrado").replace("_", " "))


def _fecha(texto, con_hora=True) -> str:
    """'2026-09-18 12:29:13' -> '18/09/2026 12:29', como se escribe en Colombia."""
    if not texto:
        return "-"
    try:
        momento = datetime.fromisoformat(str(texto))
    except ValueError:
        return str(texto)[:16]
    return momento.strftime("%d/%m/%Y %H:%M" if con_hora else "%d/%m/%Y")


def _fecha_excel(texto):
    """En la hoja la fecha va como fecha de verdad, para poder ordenar y filtrar."""
    try:
        return datetime.fromisoformat(str(texto))
    except ValueError:
        return str(texto or "")[:16]


def _rango(desde: str, hasta: str) -> str:
    if desde == hasta:
        return f"Del día {_fecha(desde, con_hora=False)}"
    return f"Del {_fecha(desde, con_hora=False)} al {_fecha(hasta, con_hora=False)}"


def _dinero(valor) -> str:
    """Formato colombiano: separador de miles con punto y sin decimales."""
    try:
        numero = float(valor or 0)
    except (TypeError, ValueError):
        numero = 0.0
    return "$ " + f"{numero:,.0f}".replace(",", ".")


def _estilos():
    hojas = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle(
            "titulo", parent=hojas["Title"], fontSize=20, leading=24,
            textColor=colors.white, alignment=0, spaceAfter=0,
        ),
        "subtitulo": ParagraphStyle(
            "subtitulo", parent=hojas["Normal"], fontSize=8, leading=12,
            textColor=colors.HexColor("#B3B3B3"),
        ),
        "seccion": ParagraphStyle(
            "seccion", parent=hojas["Normal"], fontSize=9, leading=12,
            textColor=ROJO, fontName="Helvetica-Bold", spaceAfter=4,
        ),
        "normal": ParagraphStyle(
            "normal", parent=hojas["Normal"], fontSize=9, leading=13,
        ),
        "pie": ParagraphStyle(
            "pie", parent=hojas["Normal"], fontSize=7.5, leading=11, textColor=GRIS,
        ),
    }


def _cabecera(titulo: str, subtitulo: str, ancho: float):
    """Banda negra superior con el nombre de la tienda y el título."""
    est = _estilos()
    interior = Table(
        [[
            Paragraph(f"<b>{EMPRESA['nombre'].upper()}</b>", est["titulo"]),
            Paragraph(
                f"{titulo}<br/>{subtitulo}",
                ParagraphStyle(
                    "der", parent=est["subtitulo"], alignment=2, fontSize=9,
                    textColor=colors.white, leading=13,
                ),
            ),
        ]],
        colWidths=[ancho * 0.5, ancho * 0.5],
    )
    interior.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 14),
        ("RIGHTPADDING", (0, 0), (-1, -1), 14),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
        ("BACKGROUND", (0, 0), (-1, -1), NEGRO),
        ("LINEBELOW", (0, 0), (-1, -1), 2.5, ROJO),
    ]))
    return interior


def _estilo_tabla(encabezado_hasta=0, alineacion_derecha=()):
    ordenes = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), NEGRO),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDE),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GRIS_SUAVE]),
    ]
    for col in alineacion_derecha:
        ordenes.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
    return TableStyle(ordenes)


def _pie(canvas, documento):
    """Numeración y sello de generación en cada página."""
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(GRIS)
    momento = datetime.now().strftime("%d/%m/%Y %H:%M")
    canvas.drawString(
        18 * mm, 12 * mm,
        f"{EMPRESA['nombre']} - documento generado el {momento}",
    )
    canvas.drawRightString(
        documento.pagesize[0] - 18 * mm, 12 * mm, f"Página {canvas.getPageNumber()}"
    )
    canvas.setStrokeColor(BORDE)
    canvas.line(18 * mm, 16 * mm, documento.pagesize[0] - 18 * mm, 16 * mm)
    canvas.restoreState()


# ===============================================================
# Factura en PDF
# ===============================================================
def factura_pdf(factura: dict) -> bytes:
    """Arma la factura de venta. `factura` viene de factura_dict()."""
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=14 * mm, bottomMargin=22 * mm,
        title=f"Factura {factura['numero']}",
        author=EMPRESA["nombre"],
    )
    ancho = documento.width
    est = _estilos()
    piezas = [
        _cabecera(
            f"FACTURA DE VENTA  {factura['numero']}",
            f"Emitida el {_fecha(factura.get('creado_en'))}",
            ancho,
        ),
        Spacer(1, 14),
    ]

    # --- Datos del emisor y del cliente, uno al lado del otro ---
    emisor = (
        f"<b>{EMPRESA['nombre']}</b><br/>NIT {EMPRESA['nit']}<br/>"
        f"{EMPRESA['direccion']}<br/>{EMPRESA['telefono']}<br/>{EMPRESA['email']}"
    )
    cliente = (
        f"<b>{factura['cliente_nombre']}</b><br/>"
        f"Documento: {factura.get('cliente_documento') or 'No registrado'}<br/>"
        f"{factura['cliente_email']}<br/>"
        f"{factura.get('cliente_telefono') or ''}<br/>"
        f"{factura.get('cliente_direccion') or ''}"
    )
    datos = Table(
        [
            [Paragraph("EMISOR", est["seccion"]), Paragraph("CLIENTE", est["seccion"])],
            [Paragraph(emisor, est["normal"]), Paragraph(cliente, est["normal"])],
        ],
        colWidths=[ancho * 0.5, ancho * 0.5],
    )
    datos.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (0, -1), 0),
        ("RIGHTPADDING", (1, 0), (1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    piezas += [datos, Spacer(1, 16)]

    # --- Renglones ---
    filas = [["#", "Producto", "Cant.", "Precio unitario", "Subtotal"]]
    for i, d in enumerate(factura.get("detalles", []), start=1):
        nombre = d["nombre_producto"]
        if d.get("marca_producto"):
            nombre = f"{d['marca_producto']} {nombre}"
        filas.append([
            str(i),
            Paragraph(nombre, est["normal"]),
            str(d["cantidad"]),
            _dinero(d["precio_unitario"]),
            _dinero(d["subtotal"]),
        ])
    if len(filas) == 1:
        filas.append(["", Paragraph("Sin renglones registrados", est["normal"]), "", "", ""])

    tabla = Table(
        filas,
        colWidths=[ancho * 0.06, ancho * 0.46, ancho * 0.10, ancho * 0.19, ancho * 0.19],
        repeatRows=1,
    )
    tabla.setStyle(_estilo_tabla(alineacion_derecha=(2, 3, 4)))
    piezas += [tabla, Spacer(1, 14)]

    # --- Totales, alineados a la derecha ---
    resumen = [["Subtotal", _dinero(factura["subtotal"])]]
    if float(factura.get("descuento") or 0) > 0:
        resumen.append(["Descuento", "- " + _dinero(factura["descuento"])])
    if float(factura.get("impuestos") or 0) > 0:
        resumen.append(["Impuestos", _dinero(factura["impuestos"])])
    resumen.append([
        "Envío",
        _dinero(factura["costo_envio"]) if float(factura["costo_envio"] or 0) > 0 else "Gratis",
    ])
    resumen.append(["TOTAL", _dinero(factura["total"])])

    tabla_totales = Table(resumen, colWidths=[ancho * 0.22, ancho * 0.22])
    tabla_totales.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEABOVE", (0, -1), (-1, -1), 1.2, NEGRO),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, -1), (-1, -1), 11),
        ("TEXTCOLOR", (1, -1), (1, -1), ROJO),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, BORDE),
    ]))
    contenedor = Table([["", tabla_totales]], colWidths=[ancho * 0.56, ancho * 0.44])
    contenedor.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
    ]))
    piezas += [contenedor, Spacer(1, 18)]

    # --- Estado y notas ---
    etiquetas = {"emitida": "EMITIDA", "pagada": "PAGADA", "anulada": "ANULADA"}
    nota = (
        f"<b>Estado de la factura:</b> {etiquetas.get(factura['estado'], factura['estado'])}"
        f" &nbsp;&nbsp;|&nbsp;&nbsp; <b>Pedido:</b> {factura.get('codigo_venta') or '-'}"
        f" &nbsp;&nbsp;|&nbsp;&nbsp; <b>Medio de pago:</b> "
        f"{_metodo_pago(factura.get('metodo_pago'))}"
    )
    piezas.append(Paragraph(nota, est["normal"]))
    if factura.get("observaciones"):
        piezas += [Spacer(1, 6), Paragraph(factura["observaciones"], est["pie"])]

    piezas += [
        Spacer(1, 16),
        Paragraph(
            "Documento generado electrónicamente por el sistema de PhoneStore. "
            "Este archivo corresponde a un proyecto académico de formación (SENA, "
            "ficha 3406211) y no constituye un documento fiscal válido.",
            est["pie"],
        ),
    ]

    documento.build(piezas, onFirstPage=_pie, onLaterPages=_pie)
    return buffer.getvalue()


# ===============================================================
# Reporte de ventas en PDF
# ===============================================================
def reporte_ventas_pdf(ventas: list, desde: str, hasta: str, resumen: dict) -> bytes:
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=12 * mm, bottomMargin=22 * mm,
        title=f"Reporte de ventas {desde} a {hasta}",
        author=EMPRESA["nombre"],
    )
    ancho = documento.width
    est = _estilos()
    piezas = [
        _cabecera("REPORTE DE VENTAS", _rango(desde, hasta), ancho),
        Spacer(1, 12),
    ]

    # --- Cifras del periodo ---
    tarjetas = [
        ["Ventas registradas", "Unidades vendidas", "Ingresos del período", "Ticket promedio"],
        [
            str(resumen.get("total_ventas", 0)),
            str(resumen.get("total_articulos", 0)),
            _dinero(resumen.get("ingresos", 0)),
            _dinero(resumen.get("ticket_promedio", 0)),
        ],
    ]
    tabla_tarjetas = Table(tarjetas, colWidths=[ancho / 4.0] * 4)
    tabla_tarjetas.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, 0), 7.5),
        ("TEXTCOLOR", (0, 0), (-1, 0), GRIS),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 15),
        ("TEXTCOLOR", (0, 1), (-1, 1), NEGRO),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("BOX", (0, 0), (-1, -1), 0.5, BORDE),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDE),
        ("BACKGROUND", (0, 0), (-1, -1), GRIS_SUAVE),
    ]))
    piezas += [tabla_tarjetas, Spacer(1, 14)]

    # --- Renglón por venta ---
    filas = [["Fecha", "Número", "Cliente", "Productos y servicios", "Cant.", "Total", "Estado"]]
    for v in ventas:
        articulos = "; ".join(
            f"{d['cantidad']} x {d['nombre_producto']}" for d in v.get("detalles", [])
        ) or "-"
        filas.append([
            _fecha(v.get("creado_en")),
            v["codigo"],
            Paragraph(v["cliente_nombre"], est["normal"]),
            Paragraph(articulos, est["normal"]),
            str(v.get("total_articulos", 0)),
            _dinero(v["total"]),
            v["estado"].capitalize(),
        ])
    if len(filas) == 1:
        filas.append(["", "", Paragraph("Sin ventas en el rango elegido", est["normal"]),
                      "", "", "", ""])

    tabla = Table(
        filas,
        colWidths=[
            ancho * 0.11, ancho * 0.11, ancho * 0.17,
            ancho * 0.33, ancho * 0.06, ancho * 0.12, ancho * 0.10,
        ],
        repeatRows=1,
    )
    tabla.setStyle(_estilo_tabla(alineacion_derecha=(4, 5)))
    piezas += [tabla, Spacer(1, 12)]

    cantidad = resumen.get("total_ventas", 0)
    piezas.append(Paragraph(
        f"<b>Total del período: {_dinero(resumen.get('ingresos', 0))}</b> "
        f"en {cantidad} {'venta' if cantidad == 1 else 'ventas'}. "
        "Los pedidos cancelados no suman a los ingresos.",
        est["normal"],
    ))

    documento.build(piezas, onFirstPage=_pie, onLaterPages=_pie)
    return buffer.getvalue()


# ===============================================================
# Reporte de ventas en Excel
# ===============================================================
def reporte_ventas_excel(ventas: list, desde: str, hasta: str, resumen: dict) -> bytes:
    """Hoja de cálculo con una fila por renglón de venta.

    Se desglosa producto por producto (y no una fila por venta) porque así
    la hoja sirve para filtrar y hacer tablas dinámicas, que es para lo
    que se pide el Excel.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas"

    negro = PatternFill("solid", fgColor="FF000000")
    gris = PatternFill("solid", fgColor="FFF3F4F6")
    blanco_negrita = Font(bold=True, color="FFFFFFFF", size=10)
    borde = Side(style="thin", color="FFD9D9DE")
    marco = Border(left=borde, right=borde, top=borde, bottom=borde)

    # --- Encabezado del informe ---
    hoja["A1"] = f"{EMPRESA['nombre']} - Reporte de ventas"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"] = _rango(desde, hasta)
    hoja["A2"].font = Font(size=10, color="FF6B7280")
    hoja["A3"] = f"Generado el {datetime.now():%d/%m/%Y %H:%M}"
    hoja["A3"].font = Font(size=9, color="FF6B7280")

    hoja["E1"], hoja["F1"] = "Ventas registradas", resumen.get("total_ventas", 0)
    hoja["E2"], hoja["F2"] = "Unidades vendidas", resumen.get("total_articulos", 0)
    hoja["E3"], hoja["F3"] = "Ingresos del período", float(resumen.get("ingresos", 0))
    hoja["E4"], hoja["F4"] = "Ticket promedio", float(resumen.get("ticket_promedio", 0))
    for fila in range(1, 5):
        hoja[f"E{fila}"].font = Font(bold=True, size=10)
        hoja[f"F{fila}"].alignment = Alignment(horizontal="right")
    hoja["F3"].number_format = '"$" #,##0'
    hoja["F4"].number_format = '"$" #,##0'

    # --- Tabla ---
    encabezados = [
        "Fecha", "Número de venta", "Cliente", "Correo", "Documento",
        "Producto", "Marca", "Cantidad", "Precio unitario", "Subtotal",
        "Total de la venta", "Método de pago", "Estado",
    ]
    inicio = 6
    for col, titulo in enumerate(encabezados, start=1):
        celda = hoja.cell(row=inicio, column=col, value=titulo)
        celda.fill = negro
        celda.font = blanco_negrita
        celda.border = marco
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fila = inicio + 1
    for v in ventas:
        detalles = v.get("detalles") or [{}]
        for d in detalles:
            valores = [
                _fecha_excel(v.get("creado_en")),
                v["codigo"],
                v["cliente_nombre"],
                v.get("cliente_email", ""),
                v.get("cliente_documento", ""),
                d.get("nombre_producto", "-"),
                d.get("marca_producto", ""),
                d.get("cantidad", 0),
                float(d.get("precio_unitario") or 0),
                float(d.get("subtotal") or 0),
                float(v.get("total") or 0),
                _metodo_pago(v.get("metodo_pago")),
                v["estado"].capitalize(),
            ]
            for col, valor in enumerate(valores, start=1):
                celda = hoja.cell(row=fila, column=col, value=valor)
                celda.border = marco
                celda.font = Font(size=9)
                if col == 1 and isinstance(valor, datetime):
                    celda.number_format = "dd/mm/yyyy hh:mm"
                elif col in (9, 10, 11):
                    celda.number_format = '"$" #,##0'
                    celda.alignment = Alignment(horizontal="right")
                elif col == 8:
                    celda.alignment = Alignment(horizontal="center")
            if fila % 2 == 0:
                for col in range(1, len(encabezados) + 1):
                    hoja.cell(row=fila, column=col).fill = gris
            fila += 1

    if fila == inicio + 1:
        hoja.cell(row=fila, column=1, value="Sin ventas en el rango elegido")

    # --- Acabado: filtros, anchos y panel fijo ---
    hoja.auto_filter.ref = f"A{inicio}:M{max(fila - 1, inicio)}"
    hoja.freeze_panes = f"A{inicio + 1}"
    anchos = [18, 16, 26, 30, 14, 34, 14, 10, 16, 16, 18, 16, 12]
    for col, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.row_dimensions[inicio].height = 28

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
